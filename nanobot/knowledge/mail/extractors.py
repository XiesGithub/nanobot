"""Structured, bounded extraction in a terminable spawn subprocess."""

from __future__ import annotations

import multiprocessing
import re
import time
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from nanobot.knowledge.mail.contracts import Coverage, ExtractionResult, MailStatus, TextBlock


class BodyHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines: list[tuple[str, str]] = []
        self.buffer = ""
        self.quote_depth = 0
        self.hidden_depth = 0

    def flush(self):
        if self.buffer.strip():
            self.lines.append((self.buffer.strip(), "quoted" if self.quote_depth else "current"))
        self.buffer = ""

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth += 1
        if tag == "blockquote":
            self.flush()
            self.quote_depth += 1
        if tag in {"p", "div", "br", "tr", "li", "h1", "h2", "h3"}:
            self.flush()

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth = max(0, self.hidden_depth - 1)
        if tag == "blockquote":
            self.flush()
            self.quote_depth = max(0, self.quote_depth - 1)
        if tag in {"p", "div", "tr", "li"}:
            self.flush()

    def handle_data(self, data):
        if not self.hidden_depth:
            self.buffer += data


def text_body(text: str, *, html: bool = False) -> list[TextBlock]:
    if html:
        parser = BodyHTML()
        parser.feed(text)
        parser.flush()
        return [TextBlock(ordinal=i, text=line, locator={"paragraph": i + 1, "start": 0, "end": len(line)},
                          content_role=role, role_basis="html_structure")
                for i, (line, role) in enumerate(parser.lines)]
    blocks, offset = [], 0
    quoted_tail, signature = False, False
    for number, line in enumerate(text.replace("\r\n", "\n").replace("\r", "\n").split("\n"), 1):
        if re.match(r"^\s*(?:-{2,}\s*(?:Original Message|Forwarded message|原始邮件)|Begin forwarded message:)", line, re.I):
            quoted_tail = True
        if line.strip() == "--":
            signature = True
        role, basis = "current", "plain_unquoted"
        if quoted_tail or line.lstrip().startswith(">"):
            role, basis = "quoted", "forward_separator" if quoted_tail else "quote_prefix"
        elif signature:
            role, basis = "signature", "signature_separator"
        elif re.match(r"\s*(From:|Date:|发件人[:：]|发送时间[:：]|On .+ wrote:)", line, re.I):
            role, basis = "unknown", "unverified_inline_header"
        if line.strip():
            blocks.append(TextBlock(ordinal=len(blocks), text=line,
                locator={"line": number, "start": offset, "end": offset + len(line)},
                content_role=role, role_basis=basis))
        offset += len(line) + 1
    return blocks


def check_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(e.file_size for e in entries) > 128 * 1024**2:
            raise ValueError("archive_resource_limit")
        if any(e.file_size > max(e.compress_size, 1) * 200 for e in entries):
            raise ValueError("archive_resource_limit")


def extract(path: Path, mime: str, *, body: bool, max_chars: int) -> ExtractionResult:
    blocks: list[TextBlock] = []
    coverage = Coverage()

    def add(text, locator, role="unknown", basis="attachment"):
        nonlocal blocks
        text = str(text)
        remaining = max_chars - coverage.processed_chars
        if len(text) > remaining:
            text = text[:remaining]
            coverage.complete = False
            coverage.truncated = True
            coverage.warnings.append("extraction_limit")
        if text.strip():
            blocks.append(TextBlock(ordinal=len(blocks), text=text, locator=locator,
                                    content_role=role, role_basis=basis))
            coverage.processed_chars += len(text)
            coverage.locators.append(locator)
        return not coverage.truncated

    try:
        if body or mime in {"text/plain", "text/markdown", "text/csv"}:
            # Bounded byte read also protects standalone reparse of damaged inputs.
            with path.open("rb") as stream:
                payload = stream.read(max_chars * 4 + 1)
            if len(payload) > max_chars * 4:
                coverage.complete = False
                coverage.truncated = True
                coverage.total_known = False
                coverage.warnings.append("extraction_limit")
            try:
                text = payload.decode("utf-8-sig")
            except UnicodeDecodeError:
                import chardet
                encoding = chardet.detect(payload[:65536]).get("encoding") or "utf-8"
                text = payload.decode(encoding, errors="replace")
                coverage.warnings.append("encoding_fallback:" + encoding)
            for block in text_body(text, html=mime == "text/html"):
                if not add(block.text, block.locator, block.content_role if body else "unknown",
                           block.role_basis if body else "attachment"):
                    break
        elif mime == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                return ExtractionResult(status=MailStatus.FAILED, error_code="encrypted",
                                        coverage=Coverage(complete=False))
            for page, item in enumerate(reader.pages, 1):
                text = item.extract_text() or ""
                if not text.strip():
                    coverage.complete = False
                    coverage.warnings.append(f"needs_ocr:page:{page}")
                if not add(text, {"page": page, "start": 0, "end": min(len(text), max_chars - coverage.processed_chars)}):
                    break
        elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            check_zip(path)
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            for ordinal, item in enumerate(document.iter_inner_content(), 1):
                if isinstance(item, Paragraph):
                    if not add(item.text, {"paragraph": ordinal}):
                        break
                elif isinstance(item, Table):
                    for row_number, row in enumerate(item.rows, 1):
                        if not add(" | ".join(cell.text for cell in row.cells), {"table": ordinal, "row": row_number}):
                            break
                if coverage.truncated:
                    break
        elif mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            check_zip(path)
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if not add(" | ".join(str(v) if v is not None else "" for v in row),
                                   {"sheet": sheet.title, "row": row_number}):
                            break
                    if coverage.truncated:
                        break
            finally:
                book.close()
            coverage.warnings.append("formula_cache_only")
        elif mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            check_zip(path)
            from pptx import Presentation
            for slide_number, slide in enumerate(Presentation(path).slides, 1):
                for shape_number, shape in enumerate(slide.shapes, 1):
                    text = getattr(shape, "text", "")
                    if shape.has_table:
                        text = "\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                    if not add(text, {"slide": slide_number, "shape": shape_number}):
                        break
                if coverage.truncated:
                    break
        else:
            return ExtractionResult(status=MailStatus.FAILED, error_code="unsupported_type",
                                    coverage=Coverage(complete=False))
    except Exception:
        # No parser exception text: it can contain paths, source text or passwords.
        return ExtractionResult(status=MailStatus.FAILED, error_code="parse_failed",
                                coverage=Coverage(complete=False))
    if not blocks:
        return ExtractionResult(status=MailStatus.FAILED, error_code="needs_ocr" if mime == "application/pdf" else "empty_document",
                                coverage=Coverage(complete=False, warnings=coverage.warnings))
    return ExtractionResult(blocks=blocks, coverage=coverage,
                            status=MailStatus.READY if coverage.complete else MailStatus.PARTIAL)


def _child(connection, path: str, mime: str, body: bool, max_chars: int):
    try:
        result = extract(Path(path), mime, body=body, max_chars=max_chars)
        connection.send(result.model_dump())
    finally:
        connection.close()


def extract_isolated(path: Path, mime: str, *, body: bool, max_chars: int,
                     timeout: float = 60, memory_bytes: int = 512 * 1024**2,
                     cancel_event=None) -> ExtractionResult:
    ctx = multiprocessing.get_context("spawn")
    receiving, sending = ctx.Pipe(duplex=False)
    process = ctx.Process(target=_child, args=(sending, str(path), mime, body, max_chars), daemon=True)
    process.start()
    sending.close()
    deadline = time.monotonic() + timeout
    try:
        try:
            import psutil
            monitor = psutil.Process(process.pid)
        except ImportError:
            monitor = None
        while time.monotonic() < deadline:
            if cancel_event is not None and cancel_event.is_set():
                return ExtractionResult(status=MailStatus.FAILED, error_code="parser_cancelled",
                                        coverage=Coverage(complete=False))
            if receiving.poll(0.05):
                try:
                    result = ExtractionResult.model_validate(receiving.recv())
                except EOFError:
                    break
                if monitor is None:
                    result.coverage.warnings.append("memory_monitor_unavailable")
                return result
            if monitor is not None and process.is_alive():
                try:
                    if monitor.memory_info().rss > memory_bytes:
                        return ExtractionResult(status=MailStatus.FAILED, error_code="parser_memory_limit",
                                                coverage=Coverage(complete=False))
                except psutil.NoSuchProcess:
                    pass
            if not process.is_alive():
                break
        return ExtractionResult(status=MailStatus.FAILED, error_code="parser_timeout",
                                coverage=Coverage(complete=False))
    finally:
        if process.is_alive():
            process.terminate()
        process.join(timeout=2)
        if process.is_alive():
            process.kill()
            process.join(timeout=2)
        receiving.close()
