"""Regression coverage for DOCX paragraph and table extraction."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document

import nanobot.utils.document as document_module
from nanobot.utils.document import extract_documents, extract_text


def _save_docx(path: Path, build) -> str:
    document = Document()
    build(document)
    document.save(path)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _assert_unchanged(path: Path, before_hash: str, record_property) -> None:
    after_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    record_property("fixture_sha256_before", before_hash)
    record_property("fixture_sha256_after", after_hash)
    assert after_hash == before_hash


def test_mixed_paragraphs_and_multiple_tables_follow_body_order(tmp_path, record_property):
    path = tmp_path / "mixed.docx"

    def build(document):
        document.add_paragraph("开头段落")
        first_table = document.add_table(rows=2, cols=2)
        first_table.cell(0, 0).text = "项目"
        first_table.cell(0, 1).text = "竹桥"
        first_table.cell(1, 0).text = "设备"
        first_table.cell(1, 1).text = "120"
        document.add_paragraph("两张表之间")
        second_table = document.add_table(rows=1, cols=2)
        second_table.cell(0, 0).text = "运输"
        second_table.cell(0, 1).text = "30"
        document.add_paragraph("结尾段落")

    before_hash = _save_docx(path, build)
    extracted = extract_text(path)
    assert isinstance(extracted, str)
    record_property("extracted_text", extracted)
    _assert_unchanged(path, before_hash, record_property)

    assert extracted == "开头段落\n\n项目\t竹桥\n设备\t120\n\n两张表之间\n\n运输\t30\n\n结尾段落"


def test_table_only_preserves_empty_columns_and_cell_paragraphs(tmp_path, record_property):
    path = tmp_path / "table-only.docx"

    def build(document):
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "第一行"
        table.cell(0, 0).add_paragraph("同一单元格第二段")
        table.cell(0, 2).text = "右侧"
        # Row 1 intentionally remains empty and should not become a blank output row.
        table.cell(2, 0).text = "末行"

    before_hash = _save_docx(path, build)
    extracted = extract_text(path)
    assert isinstance(extracted, str)
    record_property("extracted_text", extracted)
    _assert_unchanged(path, before_hash, record_property)

    assert extracted == "第一行\n同一单元格第二段\t\t右侧\n末行\t\t"


@pytest.mark.parametrize("with_empty_table", [False, True], ids=["empty-document", "empty-table"])
def test_empty_document_and_all_empty_table_extract_to_empty(
    tmp_path, record_property, with_empty_table
):
    path = tmp_path / "empty.docx"

    def build(document):
        if with_empty_table:
            document.add_table(rows=2, cols=3)

    before_hash = _save_docx(path, build)
    extracted = extract_text(path)
    assert extracted == ""
    record_property("extracted_text", extracted)
    _assert_unchanged(path, before_hash, record_property)


@pytest.mark.parametrize(
    ("shape", "max_length", "expected_prefix"),
    [
        ("block-separator", 3, "前\n\n"),
        ("cell-content", 5, "左侧\t值字"),
        ("row-separator", 4, "第一行\n"),
    ],
)
def test_docx_truncation_keeps_existing_prefix_and_suffix(
    tmp_path, monkeypatch, record_property, shape, max_length, expected_prefix
):
    path = tmp_path / f"truncated-{shape}.docx"

    def build(document):
        if shape == "block-separator":
            document.add_paragraph("前")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "表格"
            document.add_paragraph("后")
        elif shape == "cell-content":
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "左侧"
            table.cell(0, 1).text = "值字段"
        else:
            table = document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "第一行"
            table.cell(1, 0).text = "第二行"

    before_hash = _save_docx(path, build)
    raw_length = {
        "block-separator": len("前\n\n表格\n\n后"),
        "cell-content": len("左侧\t值字段"),
        "row-separator": len("第一行\n第二行"),
    }[shape]
    monkeypatch.setattr(document_module, "_MAX_TEXT_LENGTH", max_length)

    extracted = extract_text(path)
    expected = f"{expected_prefix}... (truncated, {raw_length} chars total)"
    assert extracted == expected
    record_property("shape", shape)
    record_property("max_text_length", max_length)
    record_property("extracted_text", extracted)
    _assert_unchanged(path, before_hash, record_property)


def test_docx_exact_length_is_not_marked_truncated(tmp_path, monkeypatch, record_property):
    path = tmp_path / "exact-length.docx"

    def build(document):
        document.add_paragraph("前")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格"

    before_hash = _save_docx(path, build)
    expected = "前\n\n表格"
    monkeypatch.setattr(document_module, "_MAX_TEXT_LENGTH", len(expected))

    extracted = extract_text(path)
    assert extracted == expected
    record_property("extracted_text", extracted)
    _assert_unchanged(path, before_hash, record_property)


def test_extract_documents_wrapper_uses_extracted_table_text(tmp_path, record_property):
    path = tmp_path / "wrapped.docx"

    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "类别"
        table.cell(0, 1).text = "金额"
        table.cell(1, 0).text = "设备"
        table.cell(1, 1).text = "120"

    before_hash = _save_docx(path, build)
    extracted, images = extract_documents("已有内容", [str(path)])
    record_property("wrapped_text", extracted)
    record_property("images", repr(images))
    _assert_unchanged(path, before_hash, record_property)

    assert extracted == "已有内容\n\n[File: wrapped.docx]\n类别\t金额\n设备\t120"
    assert images == []
