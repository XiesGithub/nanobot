import asyncio

from nanobot.config.schema import MailKnowledgeConfig
from nanobot.knowledge.mail.chunking import chunks, exact_match, terms
from nanobot.knowledge.mail.extractors import extract, extract_isolated, text_body
from nanobot.knowledge.mail.ingestion import MailWorker
from nanobot.knowledge.mail.store import MailStore
from tests.knowledge.mail.test_store import envelope


def test_current_quote_inline_and_unverified_headers():
    blocks = text_body("价格改为120元\n> 旧价格100元\n请按新价格执行\nFrom: fake@example.test\n--\n签名")
    assert [b.content_role for b in blocks] == ["current", "quoted", "current", "unknown", "signature", "signature"]
    assert blocks[1].text == "> 旧价格100元"
    html = text_body("<script>steal()</script><p>当前120元</p><blockquote><p>旧价100元</p></blockquote><p>新条款</p>", html=True)
    assert [b.content_role for b in html] == ["current", "quoted", "current"]
    assert "steal" not in str(html)


def test_tail_coverage_and_chunk_offsets(tmp_path):
    text = "正文很长。" * 50000 + "TAIL-CLAUSE付款尾款为123元"
    file = tmp_path / "long.txt"
    file.write_text(text, encoding="utf-8")
    result = extract(file, "text/plain", body=False, max_chars=2_000_000)
    assert result.coverage.complete
    segmented = list(chunks(result.blocks))
    assert any("TAIL-CLAUSE" in c["text"] for c in segmented)
    assert all(len(c["text"].encode()) <= 960 for c in segmented)
    for chunk in segmented:
        span = chunk["spans"][0]
        assert chunk["text"] == result.blocks[span["block_ordinal"]].text[span["start"]:span["end"]]
    limited = extract(file, "text/plain", body=False, max_chars=100)
    assert limited.status == "partial"
    assert limited.coverage.truncated
    assert not limited.coverage.complete


def test_docx_table_and_pdf_failure(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("First")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Order"
    table.cell(0, 1).text = "PO-001"
    doc.add_paragraph("Tail")
    path = tmp_path / "same.docx"
    doc.save(path)
    result = extract(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", body=False, max_chars=20000)
    assert [b.text for b in result.blocks] == ["First", "Order | PO-001", "Tail"]
    assert result.blocks[1].locator == {"table": 2, "row": 1}
    path.write_bytes(b"not a PDF")
    result = extract(path, "application/pdf", body=False, max_chars=20000)
    assert result.status == "failed"
    assert result.error_code == "parse_failed"


def test_scanned_and_mixed_pdf_reports_unread_pages(tmp_path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    path = tmp_path / "scan.pdf"
    writer.write(path)
    result = extract(path, "application/pdf", body=False, max_chars=10000)
    assert result.error_code == "needs_ocr"
    assert not result.coverage.complete


def test_spawn_parser_and_timeout(tmp_path):
    path = tmp_path / "document.txt"
    path.write_text("fictional content", encoding="utf-8")
    result = extract_isolated(path, "text/plain", body=False, max_chars=10000)
    assert result.blocks[0].text == "fictional content"
    result = extract_isolated(path, "text/plain", body=False, max_chars=10000, timeout=0)
    assert result.error_code == "parser_timeout"


def test_exact_numbers_and_safe_terms():
    assert exact_match("PO-001", "订单PO-001。")
    assert exact_match("PO-001", "订单 PO-001。")
    assert not exact_match("PO-001", "PO001 PO-0010")
    assert exact_match("PO-001", "po-001")
    assert "报销" in terms('"报销" OR *')


def test_worker_publishes_text_and_moves_input_without_duplicate(tmp_path):
    async def run():
        store = MailStore(tmp_path / "mail", MailKnowledgeConfig(semantic_enabled=False))
        store.initialize()
        receipt = store.accept(envelope())
        worker = MailWorker(store)
        await worker.process(store.claim())
        with store.connect() as db:
            assert db.execute("SELECT status FROM mail_items").fetchone()[0] == "ready"
            assert db.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] == 1
            assert store.pending_bytes(db) == 0
            ref = db.execute("SELECT object_ref FROM documents").fetchone()[0]
        assert store.safe_path(ref).is_file()
        assert not list((store.root / "spool").iterdir())
        store.remove(receipt.mail_id, 1)
        assert not store.safe_path(ref).exists()
    asyncio.run(run())
