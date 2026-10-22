"""Prospective synthetic gold; all names, companies and terms are fictional.

These agent-authored labels require human review before release acceptance.
The held-out split is fixed before any embedding evaluation is run.
"""

from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from nanobot.knowledge.mail.contracts import ArchiveEnvelope, ArchivePart, SourceIdentity

VERSION = "fictional-mail-v1"
ACCOUNT = UUID("bca5104e-e445-41da-a345-7cb88243f9d5")

# subject, standard supporting text, paraphrased question
CASES = [
    ("青穗培训报销", "青穗培训的交通费须附行程单及电子发票，单人额度不得超过800元。", "参加青穗课程后，车旅支出报账需要什么凭证？"),
    ("澄海采购 PO-001", "澄海订单PO-001的预付款为合同金额的30%，余款在验收合格后付清。", "澄海这笔采购先要支付多大比例的订金？"),
    ("星桥入职安排", "星桥新同事请于9月28日携带身份证原件到星河楼三层办理报到。", "星桥录用的人第一天去哪里办手续？"),
    ("松岚账户安全", "松岚系统的临时口令将在24小时后失效，首次登录必须重设密码。", "松岚给的初始登录凭据能保持有效多久？"),
    ("沧沅设备保修", "沧沅检测仪享有自签收日起18个月免费维修，但人为损坏不在范围内。", "沧沅仪器出现故障，免费维修覆盖多长时间？"),
    ("云汀会议调整", "云汀评审会改到周五下午3点，以视频会议形式召开，原线下会议室已取消。", "云汀评审现在需要到办公室现场参加吗？"),
    ("凌川差旅住宿", "凌川项目出差的酒店住宿标准为每晚450元，超额部分由个人承担。", "为凌川出差订房时，一晚预算上限是多少？"),
    ("秋禾数据交付", "秋禾实验数据须以CSV格式交付，编码统一为UTF-8，并附字段说明表。", "秋禾要求用什么文件类型和字符编码移交试验记录？"),
    ("晴浦合同终止", "晴浦服务合同可提前30日书面通知终止，已履行部分按实际工作量结算。", "停止晴浦服务前要预留多久通知对方？"),
    ("远竹门禁申请", "远竹研发楼访客需在到访前一天提交姓名和身份证后四位，由接待人审批。", "去远竹研发楼拜访，需要提前提供哪些个人资料？"),
    ("栖舟采购包装", "栖舟玻璃样品必须使用双层泡沫箱包装，外箱需张贴易碎标签。", "栖舟的玻璃件发货时怎样防止运输破损？"),
    ("青霄仓库提货", "青霄仓库仅在工作日上午9点至11点办理提货，司机须出示提货码。", "青霄仓库什么时候能取货，司机要带什么凭据？"),
    ("落杉期刊校样", "落杉期刊要求作者在收到校样后72小时内返回修改意见，逾期按原样付印。", "落杉的排版稿收到后有多长时间可以纠错？"),
    ("桃汐活动退票", "桃汐论坛开场7日前申请退票可全额退款，此后不再受理。", "桃汐论坛不去了，最迟提前几天能把报名费拿回来？"),
    ("珀宁设备校准", "珀宁温度探头每6个月需送第三方实验室校准一次，证书随设备归档。", "珀宁测温元件的外部检定周期是多久？"),
    ("明藤研究样本", "明藤样本运输全程保持2至8摄氏度，超过8小时未入库必须重新检测。", "明藤样品送检路上需要怎样的冷链条件？"),
    ("雪樟论文报送", "雪樟学术委员会只接收匿名稿件，请删除作者姓名、单位及致谢信息。", "投给雪樟委员会的文章，哪些身份信息要去掉？"),
    ("鸣溪系统维护", "鸣溪业务系统周日凌晨1点至4点维护，期间暂停在线支付，查询服务正常。", "鸣溪停机升级时还能查记录和付款吗？"),
    ("葭湖培训考核", "葭湖培训结业须同时满足出勤率90%和最终测验80分两个条件。", "通过葭湖课程拿到结业证明，要达到哪些考核指标？"),
    ("若榆订单取消", "若榆订单PO-019在生产排程确认前可免费取消，确认后收取20%材料费。", "若榆采购已经安排生产，再撤单有什么费用？"),
    ("灯榭软件许可", "灯榭软件许可可同时激活3台设备，不允许向关联公司转让账号。", "灯榭授权能装到几台机器，账号能转给子公司吗？"),
    ("雾棠合同签署", "雾棠合同需要双方授权代表签字并加盖公章，扫描件不作为最终生效版本。", "雾棠协议怎样签才正式有效，电子扫描件算不算？"),
    ("竹砚采购验收", "竹砚样机验收必须完成连续48小时稳定性测试，并提交故障日志。", "竹砚的试制机器交付前要做多久耐久运行？"),
    ("稷湾供货排期", "稷湾订单的交期从收到预付款起计算，为15个工作日，不含法定节假日。", "稷湾供货周期从哪一天开始算，节假日计入吗？"),
    ("南穹材料寄送", "南穹申请材料应邮寄纸质盖章原件至文成路16号，快递单注明项目编号。", "南穹要的文件能只交电子版吗，纸件送到哪里？"),
    ("白蓼志愿服务", "白蓼志愿者应穿深色长裤及防滑运动鞋，现场不提供更衣设施。", "参加白蓼志愿活动对衣着有什么要求？"),
    ("霁澜数据保存", "霁澜原始测量记录需保存5年，汇总报表不能替代原始文件。", "霁澜测量资料能只留下统计结果吗，底稿要留多久？"),
    ("榕渚账期调整", "榕渚新账期为发票验真后45日付款，旧的月结60日条款不再适用。", "榕渚现在核对发票后多久结款？"),
    ("桐陌展品撤场", "桐陌展会撤场安排在闭幕当日晚上6点以后，提前搬离须获得书面许可。", "桐陌展览结束前可以直接把展品搬走吗？"),
    ("水穗紧急联络", "水穗测试异常请先冻结任务队列，再联系值班工程师；不得自行重启数据库。", "水穗测试出故障时应先做什么，能直接重启数据库吗？"),
]


def corpus():
    """120 mails, 24 attachments, historical roles and partial extraction cases."""
    mails, labels = [], []
    for index in range(120):
        uid = index + 1
        if index < len(CASES):
            subject, fact, question = CASES[index]
            body = fact
            labels.append({"id": f"semantic-{uid:03}", "uid": uid, "question": question,
                           "gold_text": fact, "split": "tuning" if index % 5 == 0 else "held_out"})
        else:
            subject = f"虚构日常通知 {uid}"
            body = f"这是第{uid}份例行排班通知。工位安排维持不变，请按排班表值班，餐厅午餐正常供应。"
        parts = []
        if index < 20:
            parts.append(ArchivePart(part_id="2", mime_section="2", filename="说明.txt", mime="text/plain",
                                     content=("本附件为虚构测试资料。\n" + body).encode()))
        if index == 31:
            body = "本次报价调整为120元。\n> 上次报价100元。\n请按新报价执行。"
        if index == 32:
            body = "---------- Forwarded message ---------\nFrom: prior@example.test\nDate: 2025-01-01\n旧报价100元。"
        if index == 33:
            body = "新交期是10月2日。\n> 能否9月30日送到？\n无法提前，仍按10月2日。\nFrom: forged@example.test\n无可验证边界的旧回复。"
        if index == 34:
            subject, body = "冲突订单 CF-080", "CF-080运费由买方承担。"
        if index == 35:
            subject, body = "冲突订单 CF-080", "CF-080运费由卖方承担，审批状态待确认。"
        if index in {36, 37}:
            from docx import Document
            doc, stream = Document(), BytesIO()
            doc.add_paragraph("虚构 DOCX 示例")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "PO-001" if index == 36 else "PO-0010"
            table.cell(0, 1).text = "尾款123元" if index == 36 else "尾款999元"
            doc.save(stream)
            parts.append(ArchivePart(part_id="2", filename="合同.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", content=stream.getvalue()))
        if index in {38, 39}:
            from pypdf import PdfWriter
            from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
            writer, stream = PdfWriter(), BytesIO()
            for page_no in range(1 if index == 38 else 36):
                page = writer.add_blank_page(width=612, height=792)
                if index == 39:
                    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
                    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
                    content = DecodedStreamObject()
                    line = "TAIL-CLAUSE: Final installment is 123 yuan." if page_no == 35 else f"Fictional background page {page_no + 1}."
                    content.set_data(f"BT /F1 12 Tf 50 700 Td ({line}) Tj ET".encode())
                    page[NameObject("/Contents")] = writer._add_object(content)
            writer.write(stream)
            parts.append(ArchivePart(part_id="2", filename="扫描件.pdf" if index == 38 else "长合同.pdf", mime="application/pdf", content=stream.getvalue()))
        received = None if index in {40, 41} else datetime(2026, 8 if index % 2 else 9, 1 + index % 27, tzinfo=UTC)
        mails.append(ArchiveEnvelope(source=SourceIdentity(account_id=ACCOUNT, mailbox="INBOX", uidvalidity=7, uid=uid),
                     source_fingerprint="fictional-evaluation-only", sender_address=f"sender{uid}@example.test",
                     sender_name=f"虚构联系人{uid}", subject=subject, body=body, parts=parts, received_at=received,
                     received_time_basis="internaldate" if received else "unknown"))
    return mails, labels
